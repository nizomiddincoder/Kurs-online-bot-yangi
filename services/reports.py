"""Statistika, Excel/CSV eksport, issiq lidlar."""
import csv
import io
from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import exists, func, select
from sqlalchemy.ext.asyncio import AsyncSession

import texts
from config import Config
from db.models import Enrollment, Group, Payment, Subscription, User
from db.repo import count_pending_payments
from services.common import esc, fmt_money, now_utc, to_local


async def active_subs_count(session: AsyncSession) -> int:
    return await session.scalar(
        select(func.count()).select_from(Subscription).where(
            Subscription.status == "active", Subscription.end_at > now_utc()
        )
    ) or 0


async def build_stats(session: AsyncSession, config: Config) -> str:
    total = await session.scalar(select(func.count()).select_from(User)) or 0
    registered = await session.scalar(select(func.count()).select_from(User).where(User.registered.is_(True))) or 0
    blocked = await session.scalar(select(func.count()).select_from(User).where(User.is_blocked.is_(True))) or 0

    seg_rows = (await session.execute(select(User.segment, func.count()).group_by(User.segment))).all()
    segments = "\n".join(
        f"• {texts.SEGMENTS.get(seg or '', 'Tanlanmagan')}: <b>{cnt}</b>" for seg, cnt in seg_rows
    ) or texts.NO_DATA

    src_rows = (await session.execute(
        select(User.source, func.count()).group_by(User.source).order_by(func.count().desc()).limit(20)
    )).all()
    sources = "\n".join(f"• {esc(src) or 'direct'}: <b>{cnt}</b>" for src, cnt in src_rows) or texts.NO_DATA

    pay_rows = (await session.execute(
        select(Payment.product, func.count(), func.coalesce(func.sum(Payment.amount), 0))
        .where(Payment.status == "approved").group_by(Payment.product)
    )).all()
    payments = "\n".join(
        f"• {texts.PRODUCT_TITLES.get(p, p)}: <b>{c}</b> ta — {fmt_money(s)} so'm" for p, c, s in pay_rows
    ) or texts.NO_DATA

    from services.funnel import fmt_when, premium_left  # aylana importdan qochish

    now = now_utc()
    webinar = texts.STATS_WEBINAR.format(
        when=fmt_when(config.webinar_at),
        codeword=await session.scalar(select(func.count()).select_from(User).where(User.webinar_watched.is_(True))) or 0,
        offers=await session.scalar(select(func.count()).select_from(User).where(User.offer_expires_at > now)) or 0,
        premium_left=await premium_left(session, config), premium_total=config.premium_seats,
    )

    return texts.STATS.format(
        total=total, registered=registered, blocked=blocked,
        segments=segments, sources=sources, payments=payments, webinar=webinar,
        pay_count=sum(r[1] for r in pay_rows), pay_sum=fmt_money(sum(r[2] for r in pay_rows)),
        pending=await count_pending_payments(session),
        active_subs=await active_subs_count(session),
    )


async def hot_leads(session: AsyncSession) -> list[User]:
    """Biznes egasi + kod so'z yozgan + kurs uchun to'lamagan."""
    from services.funnel import COURSE_PRODUCTS

    paid = exists().where(Payment.user_id == User.id, Payment.product.in_(COURSE_PRODUCTS),
                          Payment.status == "approved")
    return list(await session.scalars(
        select(User).where(User.segment == "biz", User.webinar_watched.is_(True), ~paid)
        .order_by(User.codeword_at.desc())
    ))


async def build_leads_csv(session: AsyncSession) -> tuple[bytes, int]:
    """Barcha foydalanuvchilar CSV (Excel to'g'ri ochishi uchun UTF-8 BOM va «;» ajratgich)."""
    users = list(await session.scalars(select(User).order_by(User.id)))
    subs = {s.user_id: s for s in (await session.scalars(select(Subscription))).unique()}
    purchases, totals = defaultdict(list), defaultdict(int)
    for p in await session.scalars(select(Payment).where(Payment.status == "approved").order_by(Payment.id)):
        purchases[p.user_id].append(texts.PRODUCT_TITLES.get(p.product, p.product))
        totals[p.user_id] += p.amount

    def dt(value):
        return to_local(value).strftime("%Y-%m-%d %H:%M") if value else ""

    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(["ID", "Telegram ID", "Ism", "Telefon", "Username", "Segment", "Muammo", "Manba",
                     "Ro'yxatdan o'tgan", "Ro'yxat yakunlangan", "Bepul kanalda", "Kod so'z", "Kod so'z vaqti",
                     "Taklif tugashi", "Xaridlar", "Jami to'lov", "Pullik kanal gacha", "Bloklagan"])
    for u in users:
        sub = subs.get(u.id)
        writer.writerow([
            u.id, u.tg_id, u.full_name or "", u.phone or "", f"@{u.username}" if u.username else "",
            texts.SEGMENTS.get(u.segment or "", ""), (u.problem or "").replace("\n", " "), u.source or "direct",
            dt(u.created_at), "ha" if u.registered else "yo'q",
            {True: "ha", False: "yo'q"}.get(u.free_member, ""), "ha" if u.webinar_watched else "yo'q",
            dt(u.codeword_at), dt(u.offer_expires_at), ", ".join(purchases[u.id]), totals[u.id],
            dt(sub.end_at) if sub and sub.status == "active" else "", "ha" if u.is_blocked else "",
        ])
    return buf.getvalue().encode("utf-8-sig"), len(users)


def _local(dt):
    return to_local(dt).replace(tzinfo=None) if dt else None


async def build_leads_xlsx(session: AsyncSession) -> tuple[bytes, int]:
    users = list(await session.scalars(select(User).order_by(User.id)))
    subs = {s.user_id: s for s in (await session.scalars(select(Subscription))).unique()}
    groups = defaultdict(list)
    for e in (await session.scalars(select(Enrollment).order_by(Enrollment.id))).unique():
        groups[e.user_id].append(
            f"{texts.TARIFF_NAMES.get(e.tariff, e.tariff)}: {e.group.title if e.group else 'guruh kutilmoqda'}"
        )
    now = now_utc()
    purchases = defaultdict(list)
    totals = defaultdict(int)
    for p in await session.scalars(select(Payment).where(Payment.status == "approved").order_by(Payment.id)):
        purchases[p.user_id].append(texts.PRODUCT_TITLES.get(p.product, p.product))
        totals[p.user_id] += p.amount

    wb = Workbook()
    ws = wb.active
    ws.title = "Lidlar"
    headers = ["№", "Telegram ID", "Ism-familiya", "Telefon", "Username", "Segment", "Manba",
               "Sana", "Ro'yxat yakunlangan", "Xaridlar", "Jami to'lov (so'm)", "Botni bloklagan",
               "Bepul kanalga obuna", "Pullik kanal", "Kurs guruhi"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    yes_no = {True: "Ha", False: "Yo'q", None: "Tekshirilmagan"}
    for i, u in enumerate(users, 1):
        sub = subs.get(u.id)
        if sub is None:
            paid = "Yo'q"
        elif sub.status == "active" and sub.end_at > now:
            paid = f"Faol, {to_local(sub.end_at):%d.%m.%Y} gacha"
        else:
            paid = f"Tugagan ({to_local(sub.end_at):%d.%m.%Y})"
        ws.append([
            i, u.tg_id, u.full_name, u.phone, f"@{u.username}" if u.username else "",
            texts.SEGMENTS.get(u.segment or "", ""), u.source or "direct", _local(u.created_at),
            "Ha" if u.registered else "Yo'q", ", ".join(purchases[u.id]), totals[u.id],
            "Ha" if u.is_blocked else "", yes_no[u.free_member], paid, "; ".join(groups[u.id]),
        ])
    for col, width in zip("ABCDEFGHIJKLMNO", (5, 14, 25, 16, 18, 18, 12, 18, 12, 40, 16, 10, 14, 24, 35)):
        ws.column_dimensions[col].width = width
    for cell in ws["H"][1:]:
        cell.number_format = "DD.MM.YYYY HH:MM"
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(users)
